from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from .domain import PolicyArtifact


def _analysis(artifact: PolicyArtifact) -> dict[str, Any]:
    return artifact.analysis if isinstance(artifact.analysis, dict) else {}


def _items(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _strings(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _value(value: Any) -> str:
    if value is None or value == "":
        return "--"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _percent(value: Any) -> str:
    if not isinstance(value, (int, float)):
        return "--"
    return f"{round(value * 100)}%"


def _plain(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^#+\s*", "", text)
    text = re.sub(r"[【?](\d+)†L(\d+)(?:-L?(\d+))?[】?]", lambda match: f"[Source {match.group(1)}, lines {match.group(2)}-{match.group(3) or match.group(2)}]", text)
    text = text.replace("【", "[").replace("】", "]").replace("†", " - ")
    text = text.replace("\u2013", "-").replace("\u2014", "-").replace("\u2011", "-")
    text = text.replace("\u2018", "'").replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2192", "->").replace("\u2026", "...").replace("\u2265", ">=").replace("\u2264", "<=")
    text = text.replace("\u00b2", "2").replace("\u00b3", "3").replace("\u00a0", " ").replace("•", "-")
    return text


def _inline_markdown(value: Any) -> str:
    text = _plain(value)
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"(`{1,3})(.*?)\1", r"\2", text)
    text = re.sub(r"(\*\*|__)(.*?)\1", r"\2", text)
    text = re.sub(r"(?<!\w)(\*|_)(.*?)(\1)(?!\w)", r"\2", text)
    return re.sub(r"\s+", " ", text).strip()


def _markdown_blocks(value: Any) -> list[tuple[str, str]]:
    """Turn the common Markdown returned by the policy model into export blocks."""
    source = str(value or "")
    source = re.sub(r"[ \t]+(?=#{1,6}\s)", "\n", source)
    blocks: list[tuple[str, str]] = []
    for raw_line in source.splitlines():
        line = raw_line.strip()
        if not line:
            blocks.append(("", "space"))
            continue
        if line.startswith("```"):
            continue
        heading = re.match(r"^#{1,6}\s+(.+)$", line)
        if heading:
            blocks.append((_inline_markdown(heading.group(1)), "heading"))
            continue
        if re.fullmatch(r"\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?", line):
            continue
        if line.startswith("|") or " | " in line:
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            cells = [_inline_markdown(cell) for cell in cells if cell.strip()]
            if cells:
                blocks.append(("\x1f".join(cells), "table"))
            continue
        bullet = re.match(r"^(?:[-*+]\s+|\d+[.)]\s+)(.+)$", line)
        if bullet:
            blocks.append((f"- {_inline_markdown(bullet.group(1))}", "bullet"))
            continue
        if re.fullmatch(r"[-_*]{3,}", line):
            blocks.append(("", "space"))
            continue
        blocks.append((_inline_markdown(line), "body"))
    while blocks and blocks[0][1] == "space":
        blocks.pop(0)
    while blocks and blocks[-1][1] == "space":
        blocks.pop()
    return blocks


def _refs(item: dict[str, Any]) -> str:
    refs = item.get("evidence_refs")
    if not isinstance(refs, list) or not refs:
        return "No source reference attached"
    return ", ".join(str(ref) for ref in refs)


def policy_markdown(artifact: PolicyArtifact) -> str:
    analysis = _analysis(artifact)
    lines = [
        f"# {_plain(artifact.title)}",
        "",
        "**DC-TIM policy implementation brief**",
        "",
        f"- Category: {_plain(artifact.category)}",
        f"- Status: {_plain(artifact.status).replace('_', ' ').title()}",
        f"- Revision: {artifact.revision}",
        f"- Evidence status: {_plain(analysis.get('evidence_status', 'not assessed')).title()}",
        f"- Confidence: {_percent(analysis.get('confidence'))}",
        "",
        "## Executive summary",
        "",
        _plain(analysis.get("summary") or artifact.content),
        "",
        "## Policy statement",
        "",
        _plain(artifact.content),
        "",
        "## Decision readiness",
        "",
        f"- Feasibility ({_percent((analysis.get('feasibility') or {}).get('score'))}): {_plain((analysis.get('feasibility') or {}).get('rationale') or 'Not assessed.')}",
        f"- Likelihood ({_percent((analysis.get('likelihood') or {}).get('score'))}): {_plain((analysis.get('likelihood') or {}).get('rationale') or 'Not assessed.')}",
        "",
        "## Implementation roadmap",
        "",
    ]

    phases = _items(analysis.get("phases"))
    if phases:
        for index, phase in enumerate(phases, 1):
            lines += [
                f"### Phase {index}: {_plain(phase.get('label') or 'Unnamed phase')}",
                "",
                f"- Time horizon: {_plain(phase.get('time_horizon') or 'To be scheduled')}",
                f"- Reported progress: {_percent(phase.get('progress'))}",
            ]
            for action in _strings(phase.get("actions")):
                lines.append(f"- Action: {_plain(action)}")
            lines.append("")
    else:
        lines += ["No implementation phases were returned by the analysis.", ""]

    lines += ["## Performance indicators", ""]
    metrics = _items(analysis.get("metrics"))
    if metrics:
        lines += ["| Indicator | Baseline | Projected | Change | Unit | Sources |", "| --- | ---: | ---: | ---: | --- | --- |"]
        for metric in metrics:
            change = _value(metric.get("change_percent"))
            if change != "--":
                change = f"{change}%"
            lines.append(
                f"| {_plain(metric.get('label'))} | {_value(metric.get('baseline'))} | {_value(metric.get('projected'))} | {change} | {_plain(metric.get('unit'))} | {_refs(metric)} |"
            )
    else:
        lines.append("No numeric indicators were supported by the evidence.")
    lines.append("")

    lines += ["## Recommended actions", ""]
    recommendations = _items(analysis.get("recommendations"))
    if recommendations:
        for recommendation in recommendations:
            lines += [
                f"### {_plain(recommendation.get('title') or 'Recommendation')}",
                "",
                f"- Priority: {_plain(recommendation.get('priority') or 'not set').title()}",
                f"- Timeframe: {_plain(recommendation.get('timeframe') or 'To be scheduled')}",
                f"- Confidence: {_percent(recommendation.get('confidence'))}",
                f"- Detail: {_plain(recommendation.get('detail') or 'No detail returned.')}",
                f"- Expected impact: {_plain(recommendation.get('expected_impact') or 'Not specified.')}",
                f"- Sources: {_refs(recommendation)}",
                "",
            ]
    else:
        lines += ["No evidence-grounded recommendations were returned.", ""]

    lines += ["## Risk register", ""]
    risks = _items(analysis.get("risks"))
    if risks:
        for risk in risks:
            lines += [
                f"### {_plain(risk.get('title') or 'Risk')}",
                "",
                f"- Severity: {_plain(risk.get('severity') or 'not set').title()}",
                f"- Likelihood: {_percent(risk.get('likelihood'))} | Impact: {_percent(risk.get('impact'))}",
                f"- Detail: {_plain(risk.get('detail') or 'No detail returned.')}",
                f"- Mitigation: {_plain(risk.get('mitigation') or 'Not specified.')}",
                f"- Sources: {_refs(risk)}",
                "",
            ]
    else:
        lines += ["No evidence-grounded risks were returned.", ""]

    lines += ["## Open issues and next steps", ""]
    uncertainties = _strings(analysis.get("uncertainties"))
    next_steps = _strings(analysis.get("next_steps"))
    if uncertainties:
        lines.append("### Uncertainties")
        lines.extend(f"- {_plain(item)}" for item in uncertainties)
        lines.append("")
    if next_steps:
        lines.append("### Next steps")
        lines.extend(f"{index}. {_plain(item)}" for index, item in enumerate(next_steps, 1))
        lines.append("")
    if not uncertainties and not next_steps:
        lines += ["No open issues or next steps were returned.", ""]

    lines += ["## Evidence sources", ""]
    citations = _items(analysis.get("citations"))
    if citations:
        for index, citation in enumerate(citations, 1):
            excerpt = _plain(citation.get("text") or "")
            if len(excerpt) > 360:
                excerpt = f"{excerpt[:357].rstrip()}..."
            lines += [
                f"[{index}] {_plain(citation.get('title') or 'Workspace source')}",
                f"Source type: {_plain(citation.get('source_type') or 'document')} | Relevance: {_value(citation.get('score'))}",
                excerpt or "No excerpt available.",
                "",
            ]
    else:
        lines += ["No source references were stored with this policy analysis.", ""]

    lines += [
        "## Provenance",
        "",
        f"Generated by DC-TIM from policy artifact revision {artifact.revision}.",
        f"Analysis provider: {_plain(artifact.analysis_provider or 'not recorded')}.",
        f"Trace: {_plain(artifact.analysis_trace_id or 'not recorded')}.",
    ]
    return "\n".join(lines) + "\n"


def policy_docx(artifact: PolicyArtifact) -> bytes:
    from docx import Document

    analysis = _analysis(artifact)
    document = Document()
    document.add_heading(_plain(artifact.title), 0)
    document.add_paragraph(
        f"DC-TIM policy implementation brief | Category: {_plain(artifact.category)} | "
        f"Status: {_plain(artifact.status).replace('_', ' ').title()} | Revision: {artifact.revision}"
    )
    document.add_heading("Executive summary", level=1)
    document.add_paragraph(_plain(analysis.get("summary") or artifact.content))
    document.add_heading("Policy statement", level=1)
    for text, kind in _markdown_blocks(artifact.content):
        if kind == "heading":
            document.add_heading(text, level=2)
        elif kind == "bullet":
            document.add_paragraph(text[2:].strip(), style="List Bullet")
        elif kind == "table":
            document.add_paragraph(text, style="List Bullet")
        elif kind == "body":
            document.add_paragraph(text)
    document.add_heading("Decision readiness", level=1)
    for label in ("feasibility", "likelihood"):
        assessment = analysis.get(label) if isinstance(analysis.get(label), dict) else {}
        document.add_paragraph(
            f"{label.title()} ({_percent(assessment.get('score'))}): "
            f"{_plain(assessment.get('rationale') or 'Not assessed.')}",
            style="List Bullet",
        )

    document.add_heading("Implementation roadmap", level=1)
    for phase in _items(analysis.get("phases")):
        document.add_heading(_plain(phase.get("label") or "Phase"), level=2)
        document.add_paragraph(
            f"Time horizon: {_plain(phase.get('time_horizon') or 'To be scheduled')} | "
            f"Reported progress: {_percent(phase.get('progress'))}"
        )
        for action in _strings(phase.get("actions")):
            document.add_paragraph(_plain(action), style="List Bullet")

    document.add_heading("Performance indicators", level=1)
    metrics = _items(analysis.get("metrics"))
    table = document.add_table(rows=1, cols=6)
    for cell, label in zip(table.rows[0].cells, ("Indicator", "Baseline", "Projected", "Change", "Unit", "Sources")):
        cell.text = label
    for metric in metrics:
        cells = table.add_row().cells
        cells[0].text = _plain(metric.get("label"))
        cells[1].text = _value(metric.get("baseline"))
        cells[2].text = _value(metric.get("projected"))
        cells[3].text = _value(metric.get("change_percent"))
        cells[4].text = _plain(metric.get("unit"))
        cells[5].text = _refs(metric)

    for section, key, detail_key in (("Recommended actions", "recommendations", "expected_impact"), ("Risk register", "risks", "mitigation")):
        document.add_heading(section, level=1)
        for item in _items(analysis.get(key)):
            document.add_paragraph(_plain(item.get("title") or section), style="List Bullet")
            document.add_paragraph(
                f"{_plain(item.get('detail') or 'No detail returned.')} "
                f"{section[:-1] if section.endswith('s') else section} response: "
                f"{_plain(item.get(detail_key) or 'Not specified.')}; Sources: {_refs(item)}"
            )

    document.add_heading("Open issues and next steps", level=1)
    for item in _strings(analysis.get("uncertainties")):
        document.add_paragraph(_plain(item), style="List Bullet")
    for item in _strings(analysis.get("next_steps")):
        document.add_paragraph(_plain(item), style="List Number")

    document.add_heading("Evidence sources", level=1)
    for index, citation in enumerate(_items(analysis.get("citations")), 1):
        document.add_paragraph(
            f"[{index}] {_plain(citation.get('title') or 'Workspace source')} - "
            f"{_plain(citation.get('text') or 'No excerpt available.')}",
            style="List Bullet",
        )
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_lines(artifact: PolicyArtifact) -> list[tuple[str, str]]:
    analysis = _analysis(artifact)
    lines: list[tuple[str, str]] = [(artifact.title, "title"), ("DC-TIM policy implementation brief", "subtitle")]
    lines += [
        (f"Category: {_plain(artifact.category)} | Status: {_plain(artifact.status).replace('_', ' ').title()} | Revision: {artifact.revision}", "meta"),
        (f"Evidence status: {_plain(analysis.get('evidence_status', 'not assessed')).title()} | Confidence: {_percent(analysis.get('confidence'))}", "meta"),
        ("", "space"),
        ("Executive summary", "heading"),
        (_plain(analysis.get("summary") or artifact.content), "body"),
        ("Policy statement", "heading"),
        ("Decision readiness", "heading"),
    ]
    statement_index = next(index for index, line in enumerate(lines) if line == ("Policy statement", "heading"))
    lines[statement_index + 1:statement_index + 1] = _markdown_blocks(artifact.content)
    for label in ("feasibility", "likelihood"):
        assessment = analysis.get(label) if isinstance(analysis.get(label), dict) else {}
        lines.append((f"{label.title()} ({_percent(assessment.get('score'))}): {_plain(assessment.get('rationale') or 'Not assessed.')}", "body"))

    lines.append(("Implementation roadmap", "heading"))
    phases = _items(analysis.get("phases"))
    if phases:
        for index, phase in enumerate(phases, 1):
            lines.append((f"Phase {index}: {_plain(phase.get('label') or 'Unnamed phase')}", "subheading"))
            lines.append((f"Time horizon: {_plain(phase.get('time_horizon') or 'To be scheduled')} | Progress: {_percent(phase.get('progress'))}", "meta"))
            lines.extend((f"- {_plain(action)}", "bullet") for action in _strings(phase.get("actions")))
    else:
        lines.append(("No implementation phases were returned by the analysis.", "body"))

    lines.append(("Performance indicators", "heading"))
    metrics = _items(analysis.get("metrics"))
    if metrics:
        lines.append(("Indicator | Baseline | Projected | Change | Unit | Sources", "table"))
        for metric in metrics:
            change = _value(metric.get("change_percent"))
            lines.append((f"{_plain(metric.get('label'))} | {_value(metric.get('baseline'))} | {_value(metric.get('projected'))} | {change} | {_plain(metric.get('unit'))} | {_refs(metric)}", "table"))
    else:
        lines.append(("No numeric indicators were supported by the evidence.", "body"))

    for heading, key, detail_key in (("Recommended actions", "recommendations", "expected_impact"), ("Risk register", "risks", "mitigation")):
        lines.append((heading, "heading"))
        entries = _items(analysis.get(key))
        if entries:
            for item in entries:
                lines += [
                    (_plain(item.get("title") or heading), "subheading"),
                    (f"Priority/severity: {_plain(item.get('priority') or item.get('severity') or 'not set').title()}", "meta"),
                    (_plain(item.get("detail") or "No detail returned."), "body"),
                    (f"Response: {_plain(item.get(detail_key) or 'Not specified.')} | Sources: {_refs(item)}", "body"),
                ]
        else:
            lines.append((f"No evidence-grounded {heading.lower()} were returned.", "body"))

    lines.append(("Open issues and next steps", "heading"))
    for item in _strings(analysis.get("uncertainties")):
        lines.append((f"Uncertainty: {_plain(item)}", "bullet"))
    for index, item in enumerate(_strings(analysis.get("next_steps")), 1):
        lines.append((f"{index}. {_plain(item)}", "bullet"))
    if not _strings(analysis.get("uncertainties")) and not _strings(analysis.get("next_steps")):
        lines.append(("No open issues or next steps were returned.", "body"))

    lines.append(("Evidence sources", "heading"))
    citations = _items(analysis.get("citations"))
    if citations:
        for index, citation in enumerate(citations, 1):
            lines += [
                (f"[{index}] {_plain(citation.get('title') or 'Workspace source')}", "subheading"),
                (_plain(citation.get("text") or "No excerpt available."), "body"),
            ]
    else:
        lines.append(("No source references were stored with this policy analysis.", "body"))
    lines += [
        ("Provenance", "heading"),
        (f"Generated by DC-TIM from policy artifact revision {artifact.revision}.", "meta"),
        (f"Analysis provider: {_plain(artifact.analysis_provider or 'not recorded')} | Trace: {_plain(artifact.analysis_trace_id or 'not recorded')}.", "meta"),
    ]
    return lines


def policy_pdf(artifact: PolicyArtifact) -> bytes:
    """Create a designed, multi-section PDF brief for a policy artifact."""
    from html import escape

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    analysis = _analysis(artifact)
    navy = colors.HexColor("#183B56")
    blue = colors.HexColor("#2F5D8A")
    blue_soft = colors.HexColor("#F7F9FC")
    teal_soft = colors.HexColor("#F7F9FC")
    amber_soft = colors.HexColor("#F7F9FC")
    rose_soft = colors.HexColor("#F7F9FC")
    text = colors.HexColor("#273444")
    muted = colors.HexColor("#667085")
    line = colors.HexColor("#D0D5DD")
    paper = colors.HexColor("#FFFFFF")

    def safe(value: Any) -> str:
        return _inline_markdown(value).replace("\n", " ")

    def para(value: Any, style: ParagraphStyle) -> Paragraph:
        return Paragraph(escape(safe(value)).replace("\n", "<br/>"), style)

    base = getSampleStyleSheet()
    title_style = ParagraphStyle("dcTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=23, leading=27, textColor=navy, alignment=TA_LEFT, spaceAfter=5)
    subtitle_style = ParagraphStyle("dcSubtitle", parent=base["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=muted)
    section_style = ParagraphStyle("dcSection", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=blue, spaceBefore=15, spaceAfter=7)
    subheading_style = ParagraphStyle("dcSubheading", parent=base["Heading3"], fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=navy, spaceBefore=4, spaceAfter=3)
    body_style = ParagraphStyle("dcBody", parent=base["BodyText"], fontName="Helvetica", fontSize=9, leading=13, textColor=text, spaceAfter=4)
    small_style = ParagraphStyle("dcSmall", parent=base["BodyText"], fontName="Helvetica", fontSize=7.5, leading=10, textColor=muted)
    label_style = ParagraphStyle("dcLabel", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=7.3, leading=9, textColor=muted, uppercase=True, spaceAfter=3)
    value_style = ParagraphStyle("dcValue", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=15, leading=17, textColor=navy)
    card_body_style = ParagraphStyle("dcCardBody", parent=body_style, fontSize=8.4, leading=11.5, spaceAfter=0)
    bullet_style = ParagraphStyle("dcBullet", parent=body_style, leftIndent=12, firstLineIndent=-8)
    table_style = ParagraphStyle("dcTable", parent=body_style, fontSize=7.5, leading=9.5, spaceAfter=0)
    table_header_style = ParagraphStyle("dcTableHeader", parent=table_style, fontName="Helvetica-Bold", textColor=colors.white)

    def section(title: str) -> list[Any]:
        return [Paragraph(escape(title), section_style)]

    def note_box(label: str, content: Any, background: colors.Color = blue_soft) -> Table:
        table = Table([[para(label, label_style), para(content, body_style)]], colWidths=[1.2 * inch, 6.0 * inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), background),
            ("BOX", (0, 0), (-1, -1), 0.7, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ]))
        return table

    def compact_box(label: str, content: Any, background: colors.Color = blue_soft) -> Table:
        table = Table([[para(label, label_style)], [para(content, card_body_style)]], colWidths=[3.6 * inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), background),
            ("BOX", (0, 0), (-1, -1), 0.7, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        return table

    def metric_card(label: str, value: str, note: str, background: colors.Color) -> Table:
        table = Table([[para(label, label_style)], [para(value, value_style)], [para(note, small_style)]], colWidths=[1.8 * inch], rowHeights=[0.22 * inch, 0.3 * inch, 0.38 * inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), background),
            ("BOX", (0, 0), (-1, -1), 0.7, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        return table

    def card(title: str, meta: str, detail: str, response: str, background: colors.Color) -> Table:
        table = Table([
            [para(title, subheading_style), para(meta, label_style)],
            [para(detail, card_body_style), para(response, card_body_style)],
        ], colWidths=[4.4 * inch, 2.8 * inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), background),
            ("BOX", (0, 0), (-1, -1), 0.7, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        return table

    story: list[Any] = []
    status = safe(analysis.get("evidence_status", "not assessed")).title()
    hero = Table([
        [
            [para("DC-TIM / POLICY IMPLEMENTATION BRIEF", label_style), para(artifact.title, title_style), para(f"{safe(artifact.category)}  |  Revision {artifact.revision}", subtitle_style)],
            [para("EVIDENCE STATUS", label_style), para(status, value_style), para(f"Confidence: {_percent(analysis.get('confidence'))}", small_style)],
        ]
        ], colWidths=[5.45 * inch, 1.75 * inch])
    hero.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), blue_soft),
        ("BACKGROUND", (1, 0), (1, -1), teal_soft if status.lower() == "sufficient" else amber_soft),
        ("BOX", (0, 0), (-1, -1), 0.9, line),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 13),
        ("RIGHTPADDING", (0, 0), (-1, -1), 13),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ]))
    story += [hero, Spacer(1, 12)]
    story += section("At a glance")
    story.append(Table([[
        metric_card("Evidence", status, "Source-backed readiness", teal_soft if status.lower() == "sufficient" else amber_soft),
        metric_card("Confidence", _percent(analysis.get("confidence")), "Analysis confidence", blue_soft),
        metric_card("Feasibility", _percent((analysis.get("feasibility") or {}).get("score")), "Can it be delivered?", blue_soft),
        metric_card("Likelihood", _percent((analysis.get("likelihood") or {}).get("score")), "Expected adoption", blue_soft),
    ]], colWidths=[1.8 * inch] * 4, hAlign="LEFT"))

    story += section("Executive summary")
    story.append(note_box("WHAT THIS BRIEF SAYS", analysis.get("summary") or artifact.content))
    story += section("Policy statement")
    policy_blocks = _markdown_blocks(artifact.content)
    block_index = 0
    while block_index < len(policy_blocks):
        block, kind = policy_blocks[block_index]
        if kind == "space":
            story.append(Spacer(1, 3))
        elif kind == "heading":
            story.append(Paragraph(escape(block), subheading_style))
        elif kind == "bullet":
            story.append(Paragraph(escape(block), bullet_style))
        elif kind == "table":
            table_blocks: list[str] = []
            while block_index < len(policy_blocks) and policy_blocks[block_index][1] == "table":
                table_blocks.append(policy_blocks[block_index][0])
                block_index += 1
            rows = [[para(cell, table_header_style if row_index == 0 else table_style) for cell in row.split("\x1f")] for row_index, row in enumerate(table_blocks)]
            column_count = max((len(row) for row in rows), default=1)
            for row in rows:
                while len(row) < column_count:
                    row.append(para("", table_style))
            widths = {2: [2.2 * inch, 5.0 * inch], 3: [1.65 * inch, 2.55 * inch, 3.0 * inch]}.get(column_count, [7.2 * inch / column_count] * column_count)
            markdown_table = Table(rows, colWidths=widths, repeatRows=1)
            markdown_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), navy),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, blue_soft]),
                ("GRID", (0, 0), (-1, -1), 0.35, line),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(markdown_table)
            continue
        else:
            story.append(Paragraph(escape(block), body_style))
        block_index += 1

    story += section("Decision readiness")
    readiness = Table([[
        compact_box("FEASIBILITY", (analysis.get("feasibility") or {}).get("rationale") or "Not assessed.", blue_soft),
        compact_box("LIKELIHOOD", (analysis.get("likelihood") or {}).get("rationale") or "Not assessed.", blue_soft),
    ]], colWidths=[3.6 * inch, 3.6 * inch])
    story.append(readiness)

    story += section("Implementation roadmap")
    phases = _items(analysis.get("phases"))
    if phases:
        for index, phase in enumerate(phases, 1):
            actions = "<br/>".join(f"- {escape(_inline_markdown(action))}" for action in _strings(phase.get("actions"))) or "No phase actions were returned."
            phase_table = Table([[para(f"PHASE {index}", label_style), para(phase.get("label") or "Unnamed phase", subheading_style)], [para(f"Time horizon: {phase.get('time_horizon') or 'To be scheduled'} / Progress: {_percent(phase.get('progress'))}", small_style), Paragraph(actions, card_body_style)]], colWidths=[1.55 * inch, 5.65 * inch])
            phase_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), blue_soft),
                ("BACKGROUND", (1, 0), (1, -1), paper),
                ("BOX", (0, 0), (-1, -1), 0.7, line),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]))
            story.append(KeepTogether([phase_table, Spacer(1, 7)]))
    else:
        story.append(note_box("ROADMAP", "No implementation phases were returned by the analysis.", colors.HexColor("#F8FAFC")))

    story += section("Performance indicators")
    metrics = _items(analysis.get("metrics"))
    if metrics:
        rows = [[para(label, table_header_style) for label in ("Indicator", "Baseline", "Projected", "Change", "Unit", "Sources")]]
        for metric in metrics:
            change = _value(metric.get("change_percent"))
            rows.append([para(metric.get("label"), table_style), para(_value(metric.get("baseline")), table_style), para(_value(metric.get("projected")), table_style), para(change, table_style), para(metric.get("unit"), table_style), para(_refs(metric), table_style)])
        indicator_table = Table(rows, colWidths=[1.65 * inch, 0.72 * inch, 0.72 * inch, 0.62 * inch, 0.68 * inch, 2.81 * inch], repeatRows=1)
        indicator_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.35, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(indicator_table)
    else:
        story.append(note_box("INDICATORS", "No numeric indicators were supported by the evidence.", colors.HexColor("#F8FAFC")))

    story += section("Recommended actions")
    recommendations = _items(analysis.get("recommendations"))
    if recommendations:
        for item in recommendations:
            story.append(KeepTogether([card(item.get("title") or "Recommendation", f"Priority: {safe(item.get('priority') or 'not set').title()}  |  {safe(item.get('timeframe') or 'To be scheduled')}", item.get("detail") or "No detail returned.", f"Expected impact: {safe(item.get('expected_impact') or 'Not specified.')} / Sources: {safe(_refs(item))}", blue_soft), Spacer(1, 7)]))
    else:
        story.append(note_box("ACTIONS", "No evidence-grounded recommendations were returned.", colors.HexColor("#F8FAFC")))

    story += section("Risk register")
    risks = _items(analysis.get("risks"))
    if risks:
        for item in risks:
            story.append(KeepTogether([card(item.get("title") or "Risk", f"Severity: {safe(item.get('severity') or 'not set').title()}  |  Likelihood: {_percent(item.get('likelihood'))}", item.get("detail") or "No detail returned.", f"Mitigation: {safe(item.get('mitigation') or 'Not specified.')} / Sources: {safe(_refs(item))}", blue_soft), Spacer(1, 7)]))
    else:
        story.append(note_box("RISKS", "No evidence-grounded risks were returned.", colors.HexColor("#F8FAFC")))

    story += section("Open issues and next steps")
    uncertainties = _strings(analysis.get("uncertainties"))
    next_steps = _strings(analysis.get("next_steps"))
    issue_rows = []
    if uncertainties:
        issue_rows.append([para("UNCERTAINTIES", label_style), Paragraph("<br/>".join(f"- {escape(_inline_markdown(item))}" for item in uncertainties), body_style)])
    if next_steps:
        issue_rows.append([para("NEXT STEPS", label_style), Paragraph("<br/>".join(f"{index}. {escape(_inline_markdown(item))}" for index, item in enumerate(next_steps, 1)), body_style)])
    story.append(note_box("OPEN ISSUES", "No open issues or next steps were returned." if not issue_rows else "Review the items below before implementation.", amber_soft))
    if issue_rows:
        issues = Table(issue_rows, colWidths=[1.25 * inch, 5.98 * inch])
        issues.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), blue_soft), ("BOX", (0, 0), (-1, -1), 0.7, line), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story.append(Spacer(1, 7))
        story.append(issues)

    story += section("Evidence sources")
    citations = _items(analysis.get("citations"))
    if citations:
        for index, citation in enumerate(citations, 1):
            excerpt = safe(citation.get("text") or "No excerpt available.")
            if len(excerpt) > 600:
                excerpt = f"{excerpt[:597].rstrip()}..."
            story.append(KeepTogether([card(f"[{index}] {citation.get('title') or 'Workspace source'}", f"{safe(citation.get('source_type') or 'document').title()}  |  Relevance: {_value(citation.get('score'))}", excerpt, "", colors.HexColor("#F8FAFC")), Spacer(1, 7)]))
    else:
        story.append(note_box("SOURCES", "No source references were stored with this policy analysis.", colors.HexColor("#F8FAFC")))

    story += section("Provenance")
    story.append(Paragraph(escape(f"Generated by DC-TIM from policy artifact revision {artifact.revision}. Analysis provider: {artifact.analysis_provider or 'not recorded'}. Trace: {artifact.analysis_trace_id or 'not recorded'}."), small_style))

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=letter, rightMargin=0.65 * inch, leftMargin=0.65 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch, title=artifact.title, author="DC-TIM")

    def draw_page(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setStrokeColor(line)
        canvas.setLineWidth(0.6)
        canvas.line(0.65 * inch, 0.42 * inch, 7.85 * inch, 0.42 * inch)
        canvas.setFillColor(muted)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(0.65 * inch, 0.25 * inch, "DC-TIM | Evidence-grounded policy implementation brief")
        canvas.drawRightString(7.85 * inch, 0.25 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=draw_page, onLaterPages=draw_page)
    return output.getvalue()
