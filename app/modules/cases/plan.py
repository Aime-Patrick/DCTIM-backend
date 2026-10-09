from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from ..indicators.domain import CaseIndicator
from ..interventions.domain import CaseIntervention
from .domain import TransformationCase


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def build_implementation_plan(
    case: TransformationCase,
    indicators: list[CaseIndicator],
    interventions: list[CaseIntervention],
    *,
    decision: str,
    timeline: str,
    notes: str,
) -> dict[str, Any]:
    """Build a reviewable plan from approved case material.

    This deliberately does not invent numeric targets or owners. The plan is a
    structured hand-off of facts already accepted in the case workspace.
    """
    previous = case.metadata.get("implementation_plan")
    previous_version = previous.get("version", 0) if isinstance(previous, dict) else 0
    now = utc_now()
    actions = [
        {
            "title": item.name,
            "owner": str(item.metadata.get("owner") or "To be assigned"),
            "timeframe": item.timeframe,
            "priority": item.priority,
            "description": item.description,
            "source_intervention_id": item.id,
        }
        for item in interventions
        if item.status in {"approved", "in_progress", "completed"}
    ]
    plan_indicators = [
        {
            "indicator_id": item.id,
            "name": item.name,
            "unit": item.unit,
            "baseline_value": item.baseline_value,
            "target_value": item.target_value,
            "current_value": item.current_value,
            "direction": item.direction,
        }
        for item in indicators
    ]
    return {
        "case_id": case.id,
        "status": "draft",
        "version": max(1, int(previous_version) + 1),
        "title": f"{case.title} implementation plan",
        "objective": case.desired_outcome,
        "decision": decision.strip() or "Decision to be confirmed after scenario comparison.",
        "actions": actions,
        "indicators": plan_indicators,
        "timeline": timeline.strip() or "To be scheduled",
        "notes": notes.strip(),
        "created_at": now.isoformat(),
        "updated_at": now.isoformat(),
        "approved_at": None,
        "approved_by": None,
    }


def monitoring_status(
    baseline: float | None,
    target: float | None,
    current: float | None,
    direction: str,
) -> tuple[float | None, str]:
    if target is None:
        return None, "no_target"
    if current is None or baseline is None:
        return None, "not_started"
    span = target - baseline
    if span == 0:
        progress = 100.0 if current == target else 0.0
    elif direction == "decrease":
        progress = ((baseline - current) / (baseline - target)) * 100
    else:
        progress = ((current - baseline) / span) * 100
    progress = max(0.0, min(100.0, progress))
    if progress >= 100:
        return progress, "on_target"
    if progress <= 25:
        return progress, "off_track"
    return progress, "in_progress"


def build_monitoring(indicators: list[CaseIndicator]) -> dict[str, Any]:
    rows = []
    statuses: list[str] = []
    for item in indicators:
        progress, status = monitoring_status(
            item.baseline_value, item.target_value, item.current_value, item.direction
        )
        statuses.append(status)
        rows.append(
            {
                "indicator_id": item.id,
                "name": item.name,
                "unit": item.unit,
                "direction": item.direction,
                "baseline_value": item.baseline_value,
                "target_value": item.target_value,
                "current_value": item.current_value,
                "progress_percent": progress,
                "status": status,
                "measurement_date": item.measurement_date,
                "source_refs": list(item.source_refs),
            }
        )
    if not rows:
        overall = "no_indicators"
    elif all(status == "on_target" for status in statuses):
        overall = "on_track"
    elif any(status == "off_track" for status in statuses):
        overall = "off_track"
    elif all(status == "not_started" for status in statuses):
        overall = "not_started"
    else:
        overall = "in_progress"
    return {
        "generated_at": utc_now(),
        "overall_status": overall,
        "indicators": rows,
    }


def plan_markdown(case: TransformationCase, plan: dict[str, Any]) -> str:
    lines = [
        f"# {plan.get('title') or case.title}",
        "",
        f"**Case:** {case.title}",
        f"**Territory:** {case.territory or 'Not specified'}",
        f"**Status:** {plan.get('status', 'draft').title()}",
        "",
        "## Objective",
        "",
        str(plan.get("objective") or case.desired_outcome),
        "",
        "## Decision",
        "",
        str(plan.get("decision") or "To be confirmed"),
        "",
        "## Actions",
        "",
    ]
    actions = plan.get("actions") or []
    if actions:
        for action in actions:
            lines.append(
                f"- **{action.get('title', 'Action')}** -- {action.get('description', '')} "
                f"(Owner: {action.get('owner', 'To be assigned')}; "
                f"Timeframe: {action.get('timeframe', 'To be scheduled')}; "
                f"Priority: {action.get('priority', 'medium')})"
            )
    else:
        lines.append("- No actions have been approved yet.")
    lines += ["", "## Indicators", ""]
    indicators = plan.get("indicators") or []
    if indicators:
        lines.append("| Indicator | Baseline | Target | Current | Direction |")
        lines.append("| --- | ---: | ---: | ---: | --- |")
        for item in indicators:
            lines.append(
                f"| {item.get('name', '')} ({item.get('unit', '')}) | "
                f"{_value(item.get('baseline_value'))} | {_value(item.get('target_value'))} | "
                f"{_value(item.get('current_value'))} | {item.get('direction', 'neutral')} |"
            )
    else:
        lines.append("No indicators have been approved yet.")
    lines += ["", "## Timeline", "", str(plan.get("timeline") or "To be scheduled"), ""]
    if plan.get("notes"):
        lines += ["## Notes", "", str(plan["notes"]), ""]
    lines += ["---", "Generated by DC-TIM from the approved settlement case workspace."]
    return "\n".join(lines) + "\n"


def _value(value: object) -> str:
    return "--" if value is None else str(value)


def plan_docx(case: TransformationCase, plan: dict[str, Any]) -> bytes:
    from io import BytesIO
    from docx import Document

    document = Document()
    document.add_heading(str(plan.get("title") or case.title), 0)
    document.add_paragraph(f"Case: {case.title} | Territory: {case.territory or 'Not specified'}")
    document.add_heading("Objective", level=1)
    document.add_paragraph(str(plan.get("objective") or case.desired_outcome))
    document.add_heading("Decision", level=1)
    document.add_paragraph(str(plan.get("decision") or "To be confirmed"))
    document.add_heading("Actions", level=1)
    for action in plan.get("actions") or []:
        document.add_paragraph(
            f"{action.get('title', 'Action')}: {action.get('description', '')} "
            f"Owner: {action.get('owner', 'To be assigned')}; "
            f"Timeframe: {action.get('timeframe', 'To be scheduled')}; "
            f"Priority: {action.get('priority', 'medium')}",
            style="List Bullet",
        )
    document.add_heading("Indicators", level=1)
    rows = plan.get("indicators") or []
    table = document.add_table(rows=1, cols=5)
    for cell, label in zip(table.rows[0].cells, ("Indicator", "Baseline", "Target", "Current", "Direction")):
        cell.text = label
    for item in rows:
        cells = table.add_row().cells
        cells[0].text = f"{item.get('name', '')} ({item.get('unit', '')})"
        cells[1].text = _value(item.get("baseline_value"))
        cells[2].text = _value(item.get("target_value"))
        cells[3].text = _value(item.get("current_value"))
        cells[4].text = str(item.get("direction", "neutral"))
    document.add_heading("Timeline", level=1)
    document.add_paragraph(str(plan.get("timeline") or "To be scheduled"))
    if plan.get("notes"):
        document.add_heading("Notes", level=1)
        document.add_paragraph(str(plan["notes"]))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def plan_pdf(case: TransformationCase, plan: dict[str, Any]) -> bytes:
    """Create a small dependency-free text PDF for reliable server exports."""
    text_lines = plan_markdown(case, plan).replace("# ", "").replace("## ", "").splitlines()
    pages: list[list[str]] = []
    for index in range(0, len(text_lines), 46):
        pages.append(text_lines[index : index + 46])
    if not pages:
        pages = [[]]
    objects: list[bytes] = []
    page_ids: list[int] = []
    for page_number, lines in enumerate(pages, 1):
        commands = ["BT", "/F1 10 Tf", "48 748 Td", "14 TL"]
        for line in lines:
            safe = str(line).encode("latin-1", "replace").decode("latin-1").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({safe[:115]}) Tj")
            commands.append("0 -1.1 TD")
        commands += [f"ET", "BT", "/F1 8 Tf", "48 28 Td", f"(DC-TIM settlement plan | Page {page_number}/{len(pages)}) Tj", "ET"]
        stream = "\n".join(commands).encode("latin-1", "replace")
        content_id = len(objects) + 1
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append(b"")
        objects[page_id - 1] = b"PAGE_PLACEHOLDER"
    pages_id = len(objects) + 1
    catalog_id = pages_id + 1
    font_id = catalog_id + 1
    for page_number, page_id in enumerate(page_ids):
        content_id = page_id - 1
        objects[page_id - 1] = f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode()
    objects.append(f"<< /Type /Pages /Kids [{' '.join(f'{item} 0 R' for item in page_ids)}] /Count {len(page_ids)} >>".encode())
    objects.append(f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, body in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)
