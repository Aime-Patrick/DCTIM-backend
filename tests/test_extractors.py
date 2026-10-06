from __future__ import annotations

import io
import sys
import types

import pytest

from app.modules.rag.infrastructure.extractors import ExtractionError, extract_document
from app.modules.rag.text import source_locations


def test_pdf_extraction_preserves_page_provenance(monkeypatch: pytest.MonkeyPatch) -> None:
    class Page:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class Reader:
        def __init__(self, stream: io.BytesIO) -> None:
            assert stream.read() == b"pdf"
            self.pages = [Page("Page one evidence"), Page("")]

    fake_pypdf = types.ModuleType("pypdf")
    fake_pypdf.PdfReader = Reader  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "pypdf", fake_pypdf)

    result = extract_document("report.pdf", b"pdf")

    assert "[Page 1]" in result.text
    assert "[Page 2]" in result.text
    assert result.metadata["pages_with_text"] == [1]
    assert result.metadata["pages_needing_ocr"] == [2]
    assert result.metadata["ocr_required"] is True


def test_pdf_without_text_fails_instead_of_silently_losing_content(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class Page:
        def extract_text(self) -> str:
            return ""

    class Reader:
        def __init__(self, stream: io.BytesIO) -> None:
            self.pages = [Page()]

    fake_pypdf = types.ModuleType("pypdf")
    fake_pypdf.PdfReader = Reader  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "pypdf", fake_pypdf)

    with pytest.raises(ExtractionError, match="OCR"):
        extract_document("scan.pdf", b"pdf")


def test_docx_extraction_keeps_headings_tables_and_order() -> None:
    from docx import Document

    document = Document()
    document.add_heading("Education outcomes", level=1)
    document.add_paragraph("The baseline is documented below.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Year"
    table.cell(0, 1).text = "Literacy"
    table.cell(1, 0).text = "2024"
    table.cell(1, 1).text = "89%"
    stream = io.BytesIO()
    document.save(stream)

    result = extract_document("report.docx", stream.getvalue())

    assert "# Education outcomes" in result.text
    assert "[DOCX table 1]" in result.text
    assert "| 2024 | 89% |" in result.text
    assert result.metadata["table_count"] == 1


def test_source_locations_are_deduplicated() -> None:
    text = "[Page 1]\nA\n[Page 1]\nB\n[DOCX table 2]\nC"
    assert source_locations(text) == ["[Page 1]", "[DOCX table 2]"]
